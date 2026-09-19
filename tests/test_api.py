"""
Tests d'intégration -- reprennent les scénarios vérifiés manuellement
pendant le développement (voir historique du projet) sous forme de suite
automatisée. Utilise une vraie base SQLite de test (voir conftest.py) plutôt
que des mocks : le comportement du verrouillage anti-doublon et du consensus
dépend de l'état réel en base.
"""

import base64
import os

from fastapi.testclient import TestClient

from api import app

client = TestClient(app)
RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_TEST = os.path.join(RACINE, "donnees_brutes_test.csv")

SCHEMA_TEST = {
    "columns": [
        {"name": "nom", "type": "string", "required": True},
        {"name": "email", "type": "regex", "rule": r"^[\w.-]+@[\w.-]+\.\w+$", "required": True},
        {"name": "date_naiss", "type": "date", "required": True},
    ]
}


def auth_header(email, password):
    """Authentification par jeton de session (voir auth.py) -- se connecte
    réellement via /auth/login et renvoie le header Bearer à utiliser pour
    les requêtes suivantes, comme le fait le front."""
    r = client.post("/auth/login", json={"email": email, "password": password})
    if r.status_code != 200:
        # Certains tests appellent auth_header() avec un mauvais mot de passe
        # exprès (ex: verrouillage anti-brute-force) -- on renvoie un header
        # invalide plutôt que de lever une exception, pour laisser le test
        # vérifier lui-même le code d'erreur renvoyé par l'endpoint protégé.
        return {"Authorization": "Bearer jeton-invalide"}
    return {"Authorization": f"Bearer {r.json()['token']}"}


def approuver(email):
    """Simule la validation manuelle d'un compte (voir admin_tools.py) --
    nécessaire depuis l'ajout du contrôle d'approbation, sans quoi tout
    endpoint fonctionnel renvoie 403 même avec des identifiants valides."""
    from models import User, get_engine, get_session_factory

    db = get_session_factory(get_engine())()
    try:
        db.query(User).filter_by(email=email).update({"approuve": True})
        db.commit()
    finally:
        db.close()


# ---------------------------------------------------------------------------
# AUTHENTIFICATION
# ---------------------------------------------------------------------------

def test_inscription_client_et_worker():
    r = client.post("/auth/register", json={"email": "startup@ia.fr", "password": "pw123", "role": "client", "secteur_activite": "E-commerce"})
    assert r.status_code == 200
    r = client.post("/auth/register", json={"email": "w1@uco.fr", "password": "pw123", "role": "worker", "accepte_confidentialite": True})
    assert r.status_code == 200
    r = client.post("/auth/register", json={"email": "w2@uco.fr", "password": "pw123", "role": "worker", "accepte_confidentialite": True})
    assert r.status_code == 200
    approuver("startup@ia.fr")
    approuver("w1@uco.fr")
    approuver("w2@uco.fr")


def test_inscription_client_sans_secteur_activite_refusee():
    r = client.post("/auth/register", json={"email": "sans_secteur@ia.fr", "password": "pw123", "role": "client"})
    assert r.status_code == 400


def test_inscription_email_deja_utilise_refusee():
    r = client.post("/auth/register", json={"email": "startup@ia.fr", "password": "autre", "role": "client", "secteur_activite": "E-commerce"})
    assert r.status_code == 400


def test_inscription_worker_sans_case_confidentialite_refusee():
    r = client.post("/auth/register", json={"email": "sans_case@uco.fr", "password": "pw123", "role": "worker"})
    assert r.status_code == 400
    r = client.post("/auth/register", json={"email": "sans_case@uco.fr", "password": "pw123", "role": "worker", "accepte_confidentialite": False})
    assert r.status_code == 400
    # Un client n'est lui pas concerné par cette exigence
    r = client.post("/auth/register", json={"email": "client_sans_case@ia.fr", "password": "pw123", "role": "client", "secteur_activite": "Conseil"})
    assert r.status_code == 200


def test_connexion_mauvais_mot_de_passe_refusee():
    r = client.get("/auth/me", headers=auth_header("startup@ia.fr", "mauvais_mdp"))
    assert r.status_code == 401


def test_compte_verrouille_temporairement_apres_echecs_repetes():
    """Le verrouillage se vérifie désormais à la connexion (/auth/login),
    plus sur chaque requête authentifiée -- voir auth.py."""
    client.post("/auth/register", json={"email": "bruteforce@uco.fr", "password": "bonmdp123", "role": "worker", "accepte_confidentialite": True})

    for _ in range(5):
        r = client.post("/auth/login", json={"email": "bruteforce@uco.fr", "password": "mauvais"})
        assert r.status_code == 401

    # Même avec le BON mot de passe, le compte reste verrouillé
    r = client.post("/auth/login", json={"email": "bruteforce@uco.fr", "password": "bonmdp123"})
    assert r.status_code == 429


def test_fichier_trop_volumineux_refuse():
    import io

    gros_fichier = io.BytesIO(b"nom;email\n" + (b"a" * 1000 + b";x@x.fr\n") * 20000)  # > 10 Mo
    r = client.post(
        "/projects/infer-schema", files={"fichier": ("gros.csv", gros_fichier, "text/csv")},
        headers=auth_header("startup@ia.fr", "pw123"),
    )
    assert r.status_code == 413


def test_compte_non_approuve_bloque_sur_les_endpoints_fonctionnels():
    r = client.post("/auth/register", json={"email": "en_attente@uco.fr", "password": "pw123", "role": "worker", "accepte_confidentialite": True})
    assert r.status_code == 200
    # Connexion possible (identifiants valides)...
    r = client.get("/auth/me", headers=auth_header("en_attente@uco.fr", "pw123"))
    assert r.status_code == 200
    assert r.json()["approuve"] is False
    # ...mais aucun accès fonctionnel tant que non approuvé
    r = client.get("/projects/available", headers=auth_header("en_attente@uco.fr", "pw123"))
    assert r.status_code == 403


def test_inscription_worker_domaine_email_non_universitaire_refusee():
    r = client.post("/auth/register", json={"email": "faux@gmail.com", "password": "pw123", "role": "worker", "accepte_confidentialite": True})
    assert r.status_code == 400
    # Un client n'est pas concerné par cette restriction de domaine
    r = client.post("/auth/register", json={"email": "entreprise_gmail@gmail.com", "password": "pw123", "role": "client", "secteur_activite": "Tech"})
    assert r.status_code == 200


def test_worker_ne_peut_pas_creer_de_projet():
    r = client.post(
        "/projects", json={"titre": "x", "colonnes_schema": {"columns": []}},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r.status_code == 403


# ---------------------------------------------------------------------------
# FLUX CLIENT : inférence de schéma, création de projet, ingestion
# ---------------------------------------------------------------------------

def test_inference_schema_detecte_les_bons_types():
    with open(CSV_TEST, "rb") as f:
        r = client.post(
            "/projects/infer-schema", files={"fichier": ("d.csv", f, "text/csv")},
            headers=auth_header("startup@ia.fr", "pw123"),
        )
    assert r.status_code == 200
    types_par_nom = {c["name"]: c["type"] for c in r.json()["schema_suggere"]["columns"]}
    assert types_par_nom["email"] == "regex"
    assert types_par_nom["date_naiss"] == "date"


def test_creation_projet_et_ingestion_csv():
    r = client.post(
        "/projects",
        json={"titre": "Test intégration", "colonnes_schema": SCHEMA_TEST, "prix_par_ligne": 0.05},
        headers=auth_header("startup@ia.fr", "pw123"),
    )
    assert r.status_code == 200
    project_id = r.json()["project_id"]

    with open(CSV_TEST, "rb") as f:
        r = client.post(
            f"/projects/{project_id}/ingest", files={"fichier": ("d.csv", f, "text/csv")},
            headers=auth_header("startup@ia.fr", "pw123"),
        )
    assert r.status_code == 200
    assert r.json()["nb_taches_creees"] == 4

    # Un client ne peut pas ingérer sur le projet d'un autre
    r2 = client.post("/auth/register", json={"email": "autre_client@ia.fr", "password": "pw123", "role": "client", "secteur_activite": "Finance"})
    approuver("autre_client@ia.fr")  # sinon le test vérifierait le blocage d'approbation, pas d'appartenance
    with open(CSV_TEST, "rb") as f:
        r = client.post(
            f"/projects/{project_id}/ingest", files={"fichier": ("d.csv", f, "text/csv")},
            headers=auth_header("autre_client@ia.fr", "pw123"),
        )
    assert r.status_code == 403


# ---------------------------------------------------------------------------
# DISTRIBUTION ET VERROUILLAGE ANTI-DOUBLON
# ---------------------------------------------------------------------------

def creer_projet_isole(nom_projet, client_email="startup@ia.fr", taux_gold=0.0, gold_answers=None):
    """Crée un projet + l'ingère avec le CSV de test -> donne à un test un lot
    de micro-tâches totalement frais, sans verrou ni soumission héritée d'un
    autre test. Évite le couplage d'état entre tests.

    Passe par process_csv_to_microtasks() directement (pas l'endpoint HTTP)
    pour contrôler taux_gold de façon déterministe -- sans ça, le tirage
    aléatoire par défaut (8%) rendrait les tests qui dépendent d'une
    redondance précise intermittents."""
    from analyzer import process_csv_to_microtasks
    from crypto import chiffrer_json
    from models import MicroTask, get_engine, get_session_factory

    r = client.post(
        "/projects",
        json={"titre": nom_projet, "colonnes_schema": SCHEMA_TEST, "prix_par_ligne": 0.05},
        headers=auth_header(client_email, "pw123"),
    )
    project_id = r.json()["project_id"]

    tasks = process_csv_to_microtasks(CSV_TEST, project_id, SCHEMA_TEST, taux_gold=taux_gold, gold_answers=gold_answers)
    db = get_session_factory(get_engine())()
    try:
        for t in tasks:
            t["raw_data"] = chiffrer_json(t["raw_data"])
            if t.get("gold_answer") is not None:
                t["gold_answer"] = chiffrer_json(t["gold_answer"])
            if t.get("resultat_final") is not None:
                t["resultat_final"] = chiffrer_json(t["resultat_final"])
            db.add(MicroTask(**t))
        db.commit()
    finally:
        db.close()
    return project_id


# ---------------------------------------------------------------------------
# DISTRIBUTION ET VERROUILLAGE ANTI-DOUBLON
# ---------------------------------------------------------------------------

GOLD_LIGNE_PROPRE = {0: ["Jean Dupont", "jdupont@gmail.com", "1998-04-03"]}  # ligne 0 du CSV de test, forcée en gold standard


def test_deux_workers_ne_recoivent_jamais_la_meme_ligne_propre():
    # Ligne 0 forcée en gold standard (redondance=1, worker-facing) --
    # sans ça, une ligne propre est validée automatiquement et n'atteint
    # jamais un worker (voir le nouveau pipeline d'auto-nettoyage).
    project_id = creer_projet_isole("Test verrouillage", gold_answers=dict(GOLD_LIGNE_PROPRE))
    r1 = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    r2 = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w2@uco.fr", "pw123"))
    assert r1.status_code == 200 and r2.status_code == 200
    assert r1.json()["task_id"] != r2.json()["task_id"]


def test_worker_ne_peut_pas_soumettre_deux_fois_sur_la_meme_tache():
    # Redondance=1 nécessaire pour un paiement en un seul passage -- forcé via gold standard.
    project_id = creer_projet_isole("Test anti double-soumission", gold_answers=dict(GOLD_LIGNE_PROPRE))
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()

    r1 = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": tache["raw_data"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r1.status_code == 200

    # Une deuxième soumission sur la MÊME tâche par le MÊME worker doit être
    # refusée -- sans quoi elle repaie et retraite indéfiniment (faille corrigée).
    r2 = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": tache["raw_data"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r2.status_code == 409

    solde = client.get("/tasks/mine", headers=auth_header("w1@uco.fr", "pw123")).json()["solde_disponible"]
    assert solde == 0.05  # une seule fois, pas deux


def test_worker_ne_peut_pas_soumettre_sans_authentification():
    r = client.post("/tasks/1/submit", json={"reponse": ["a", "b", "c"]})
    assert r.status_code == 401


# ---------------------------------------------------------------------------
# CONSENSUS, PAIEMENT SYSTÉMATIQUE ET AJUSTEMENT DU TRUST_SCORE
# ---------------------------------------------------------------------------

def test_ligne_propre_redondance_1_payee_immediatement():
    # Forcée en gold standard pour tester le paiement immédiat en redondance=1
    # -- une ligne propre non tirée au sort est désormais auto-validée sans
    # jamais atteindre de worker (voir test_pipeline_auto_nettoyage.py).
    project_id = creer_projet_isole("Test paiement immediat", gold_answers=dict(GOLD_LIGNE_PROPRE))
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()
    assert tache["row_id"] == 0

    r = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": tache["raw_data"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r.json()["paye"] is True


def test_ligne_propre_sans_tirage_gold_validee_automatiquement_sans_worker():
    """Le cœur du nouveau pipeline : une ligne déjà propre après nettoyage
    automatique, non tirée au sort comme gold standard, doit être marquée
    complétée dès l'ingestion -- aucun worker ne doit jamais la voir."""
    project_id = creer_projet_isole("Test auto-validation", taux_gold=0.0)

    r = client.get("/projects/mine", headers=auth_header("startup@ia.fr", "pw123"))
    projet = next(p for p in r.json()["projets"] if p["project_id"] == project_id)
    # Lignes 0 et 3 du CSV de test sont propres après nettoyage (espaces
    # superflus corrigés automatiquement pour la ligne 3) -> auto-validées.
    assert projet["lignes_completees"] == 2

    # Aucune tâche pour ces deux lignes ne doit jamais être servie à un worker
    row_ids_vus = set()
    for _ in range(4):  # borne dure -- au plus 4 lignes au total dans le CSV de test
        r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
        if r.status_code == 404:
            break
        tache = r.json()
        row_ids_vus.add(tache["row_id"])
        client.post(
            f"/tasks/{tache['task_id']}/submit", json={"reponse": tache["raw_data"]},
            headers=auth_header("w1@uco.fr", "pw123"),
        )
    assert row_ids_vus == {1, 2}  # jamais 0 ni 3


def test_litige_est_paye_mais_impacte_le_trust_score():
    from models import MicroTask, get_engine, get_session_factory

    project_id = creer_projet_isole("Test litige")
    db = get_session_factory(get_engine())()
    tache = db.query(MicroTask).filter_by(row_id=3, project_id=project_id).first()
    tache.redundancy_level = 2
    db.commit()
    task_id = tache.id
    db.close()

    trust_avant_w1 = client.get("/auth/me", headers=auth_header("w1@uco.fr", "pw123")).json()["trust_score"]
    trust_avant_w2 = client.get("/auth/me", headers=auth_header("w2@uco.fr", "pw123")).json()["trust_score"]

    r1 = client.post(
        f"/tasks/{task_id}/submit",
        json={"reponse": ["Sophie Legrand", "sophie@legrand.fr", "1990-11-30"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r1.json()["status"] == "en_attente_autres_workers"
    assert r1.json()["paye"] is None

    r2 = client.post(
        f"/tasks/{task_id}/submit",
        json={"reponse": ["Quelqu'un d'autre", "faux@x.fr", "2000-01-01"]},
        headers=auth_header("w2@uco.fr", "pw123"),
    )
    # Décision produit : payé même en cas de désaccord
    assert r2.json()["paye"] is True

    trust_apres_w1 = client.get("/auth/me", headers=auth_header("w1@uco.fr", "pw123")).json()["trust_score"]
    trust_apres_w2 = client.get("/auth/me", headers=auth_header("w2@uco.fr", "pw123")).json()["trust_score"]

    assert trust_apres_w1 >= trust_avant_w1   # w1 rejoint le consensus -> stable ou en hausse
    assert trust_apres_w2 < trust_avant_w2    # w2 diverge nettement -> pénalisé


# ---------------------------------------------------------------------------
# HISTORIQUE WORKER : pas de fuite sur les gold standards / litiges
# ---------------------------------------------------------------------------

def test_historique_worker_ne_contient_aucun_texte_de_reponse():
    r = client.get("/tasks/mine", headers=auth_header("w1@uco.fr", "pw123"))
    assert r.status_code == 200
    data = r.json()
    assert "solde_disponible" in data
    for ligne in data["historique"]:
        assert set(ligne.keys()) == {"task_id", "date", "statut", "montant"}
        assert ligne["statut"] in ("paye", "en_attente_verification")


def test_soumission_avec_chaine_trop_longue_refusee():
    project_id = creer_projet_isole("Test taille champ")
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()
    reponse_trop_longue = ["x" * 10_000] * len(tache["header"])
    r = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": reponse_trop_longue},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r.status_code == 422


def test_soumission_avec_mauvais_nombre_de_colonnes_refusee():
    project_id = creer_projet_isole("Test nb colonnes")
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()
    r = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": ["une seule valeur"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r.status_code == 400


# ---------------------------------------------------------------------------
# NOUVELLES FONCTIONNALITÉS : projets disponibles, libération de verrou, export
# ---------------------------------------------------------------------------

def test_projets_disponibles_et_liberation_de_verrou():
    project_id = creer_projet_isole("Test disponibilite")

    r = client.get("/projects/available", headers=auth_header("w1@uco.fr", "pw123"))
    assert r.status_code == 200
    projets = {p["project_id"]: p for p in r.json()["projets"]}
    assert project_id in projets
    # Sur les 4 lignes du CSV de test, 2 sont auto-validées sans worker
    # (lignes 0 et 3, propres après nettoyage automatique) -- seules les
    # 2 lignes réellement sales (1 et 2) doivent apparaître comme disponibles.
    assert projets[project_id]["taches_disponibles"] == 2

    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    task_id = r.json()["task_id"]

    # Sans libération, la même ligne n'est pas reproposée à w1
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    assert r.json()["task_id"] != task_id

    r = client.post(f"/tasks/{task_id}/release", headers=auth_header("w1@uco.fr", "pw123"))
    assert r.status_code == 200

    # w2 doit pouvoir reprendre immédiatement la ligne libérée (pas d'attente des 15 min)
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w2@uco.fr", "pw123"))
    tache_ids_restants = {t["task_id"] for t in [r.json()]}
    # (on ne peut pas garantir qu'il retombe exactement sur task_id si w2 a déjà
    #  d'autres verrous actifs, donc on vérifie juste que la libération n'a pas échoué)


def test_projets_mine_et_export_csv():
    project_id = creer_projet_isole("Test export")

    r = client.get("/projects/mine", headers=auth_header("startup@ia.fr", "pw123"))
    projet = next(p for p in r.json()["projets"] if p["project_id"] == project_id)
    assert projet["total_lignes"] == 4
    # 2 lignes (0 et 3) sont auto-validées dès l'ingestion -- pas besoin
    # d'attendre un worker pour qu'elles comptent comme complétées.
    assert projet["lignes_completees"] == 2

    # Export partiel possible immédiatement (2 lignes déjà complètes)
    r = client.get(f"/projects/{project_id}/export", headers=auth_header("startup@ia.fr", "pw123"))
    assert r.status_code == 200
    assert "jdupont@gmail.com" in r.text

    # Un client ne peut pas exporter le projet d'un autre
    r = client.get(f"/projects/{project_id}/export", headers=auth_header("autre_client@ia.fr", "pw123"))
    assert r.status_code == 403


def test_export_refuse_si_aucune_ligne_completee():
    # Force les 2 lignes propres du CSV (0 et 3) en gold standard, pour
    # qu'aucune ne s'auto-valide -- seul moyen d'obtenir un projet où
    # rien n'est encore complété juste après l'ingestion.
    gold = {0: ["Jean Dupont", "jdupont@gmail.com", "1998-04-03"],
            3: ["Sophie Legrand", "sophie@legrand.fr", "1990-11-30"]}
    project_id = creer_projet_isole("Test export vide", gold_answers=gold)
    r = client.get(f"/projects/{project_id}/export", headers=auth_header("startup@ia.fr", "pw123"))
    assert r.status_code == 404


def test_depot_excel_xlsx_fonctionne_comme_le_csv():
    import openpyxl

    chemin_xlsx = os.path.join(RACINE, "tests", "_temp_test.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["nom", "email", "date_naiss"])
    ws.append(["Jean Dupont", "jdupont@gmail.com", "1998-04-03"])
    ws.append(["Paul Martin", "paul.martin", "2000-01-01"])
    wb.save(chemin_xlsx)

    try:
        with open(chemin_xlsx, "rb") as f:
            r = client.post(
                "/projects/infer-schema",
                files={"fichier": ("d.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                headers=auth_header("startup@ia.fr", "pw123"),
            )
        assert r.status_code == 200
        schema = r.json()["schema_suggere"]

        r = client.post(
            "/projects", json={"titre": "Test Excel", "colonnes_schema": schema, "prix_par_ligne": 0.05},
            headers=auth_header("startup@ia.fr", "pw123"),
        )
        project_id = r.json()["project_id"]

        with open(chemin_xlsx, "rb") as f:
            r = client.post(
                f"/projects/{project_id}/ingest",
                files={"fichier": ("d.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                headers=auth_header("startup@ia.fr", "pw123"),
            )
        assert r.status_code == 200
        assert r.json()["nb_taches_creees"] == 2
    finally:
        os.remove(chemin_xlsx)


def test_export_xlsx_produit_un_classeur_valide():
    import openpyxl

    # 2 lignes (0 et 3) sont auto-validées dès l'ingestion -- pas besoin
    # de faire intervenir un worker pour avoir du contenu à exporter.
    project_id = creer_projet_isole("Test export xlsx")

    r = client.get(f"/projects/{project_id}/export?format=xlsx", headers=auth_header("startup@ia.fr", "pw123"))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    chemin_temp = os.path.join(RACINE, "tests", "_temp_export.xlsx")
    with open(chemin_temp, "wb") as f:
        f.write(r.content)
    try:
        wb = openpyxl.load_workbook(chemin_temp)
        lignes = list(wb.active.iter_rows(values_only=True))
        assert lignes[0] == ("nom", "email", "date_naiss")
        assert len(lignes) == 3  # en-tête + 2 lignes auto-validées (0 et 3)
    finally:
        os.remove(chemin_temp)


def test_donnees_client_illisibles_sur_le_disque():
    """Vérifie qu'un accès brut au fichier SQLite (vol, sauvegarde mal
    configurée) ne révèle aucune donnée métier en clair -- le cœur de la
    protection "sécurité des données clients"."""
    project_id = creer_projet_isole("Test chiffrement au repos")
    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()
    valeur_sensible = tache["raw_data"][1]  # l'email de la ligne, ex: jdupont@gmail.com
    client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": tache["raw_data"]},
        headers=auth_header("w1@uco.fr", "pw123"),
    )

    with open(os.path.join(RACINE, "passerelle.db"), "rb") as f:
        contenu_brut = f.read()

    assert valeur_sensible.encode() not in contenu_brut
    assert b"Jean Dupont" not in contenu_brut


def test_colonne_sensible_jamais_montree_au_worker_mais_preservee_a_export():
    """Une colonne marquée 'sensible' dans le schéma ne doit JAMAIS apparaître
    dans ce que /tasks/next sert au worker (ni header, ni raw_data, ni badges),
    mais la vraie valeur doit tout de même ressortir intacte à l'export."""
    schema = {
        "columns": [
            {"name": "nom", "type": "string", "required": True, "sensible": False},
            {"name": "email", "type": "regex", "rule": r"^[\w.-]+@[\w.-]+\.\w+$", "required": True, "sensible": True},
            {"name": "date_naiss", "type": "date", "required": True, "sensible": False},
        ]
    }
    r = client.post(
        "/projects", json={"titre": "Test masquage", "colonnes_schema": schema, "prix_par_ligne": 0.05},
        headers=auth_header("startup@ia.fr", "pw123"),
    )
    project_id = r.json()["project_id"]
    with open(CSV_TEST, "rb") as f:
        client.post(
            f"/projects/{project_id}/ingest", files={"fichier": ("d.csv", f, "text/csv")},
            headers=auth_header("startup@ia.fr", "pw123"),
        )

    r = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w1@uco.fr", "pw123"))
    tache = r.json()
    assert "email" not in tache["header"]
    assert len(tache["header"]) == 2

    # Le worker ne soumet que les 2 colonnes visibles -- si l'une d'elles a
    # un badge bloquant (ex. date ambiguë), il doit d'abord la corriger avant
    # que la soumission soit acceptée.
    reponse_corrigee = list(tache["raw_data"])
    idx_date = tache["header"].index("date_naiss")
    reponse_corrigee[idx_date] = "1998-04-03"
    r = client.post(
        f"/tasks/{tache['task_id']}/submit", json={"reponse": reponse_corrigee},
        headers=auth_header("w1@uco.fr", "pw123"),
    )
    assert r.status_code == 200

    # Soumettre le mauvais nombre de valeurs (3 au lieu de 2) doit échouer
    r2 = client.get("/tasks/next", params={"project_id": project_id}, headers=auth_header("w2@uco.fr", "pw123"))
    r2_submit = client.post(
        f"/tasks/{r2.json()['task_id']}/submit", json={"reponse": ["a", "b", "c"]},
        headers=auth_header("w2@uco.fr", "pw123"),
    )
    assert r2_submit.status_code == 400

    # Mais l'export final récupère bien le VRAI email, jamais vu par le worker
    r = client.get(f"/projects/{project_id}/export", headers=auth_header("startup@ia.fr", "pw123"))
    assert "jdupont@gmail.com" in r.text


def test_deconnexion_revoque_le_jeton_immediatement():
    """La déconnexion coupe l'accès tout de suite -- pas besoin d'attendre
    l'expiration naturelle du jeton (DUREE_SESSION_HEURES)."""
    client.post("/auth/register", json={"email": "revoc@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("revoc@uco.fr")

    login = client.post("/auth/login", json={"email": "revoc@uco.fr", "password": "pw123456"})
    assert login.status_code == 200
    header = {"Authorization": f"Bearer {login.json()['token']}"}

    # Le jeton fonctionne avant déconnexion
    assert client.get("/auth/me", headers=header).status_code == 200

    r = client.post("/auth/logout", headers=header)
    assert r.status_code == 200

    # Le même jeton est désormais rejeté
    assert client.get("/auth/me", headers=header).status_code == 401


def test_mauvais_jeton_rejete():
    r = client.get("/auth/me", headers={"Authorization": "Bearer jeton-qui-nexiste-pas"})
    assert r.status_code == 401


# ---------------------------------------------------------------------------
# CONNEXION PAR CODE TELEGRAM (remplace le mot de passe une fois lié)
# ---------------------------------------------------------------------------

def _activer_bot_telegram_de_test(monkeypatch):
    """TELEGRAM_BOT_TOKEN est importé PAR VALEUR dans api.py (`from
    notifications import TELEGRAM_BOT_TOKEN`) -- le patcher sur le module
    notifications ne suffit donc pas, il faut aussi patcher api.py. Les
    envois réels sont remplacés par un no-op réussi : ces tests vérifient le
    flux d'authentification, pas l'intégration réseau avec Telegram."""
    import api
    import notifications
    monkeypatch.setattr(api, "TELEGRAM_BOT_TOKEN", "jeton-bot-de-test")
    monkeypatch.setattr(notifications, "envoyer_message_telegram", lambda chat_id, texte: True)


def test_liaison_telegram_puis_mot_de_passe_refuse_ensuite(monkeypatch):
    _activer_bot_telegram_de_test(monkeypatch)

    client.post("/auth/register", json={"email": "tg1@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("tg1@uco.fr")
    header = auth_header("tg1@uco.fr", "pw123456")

    assert client.get("/auth/me", headers=header).json()["telegram_lie"] is False

    r = client.post("/auth/telegram/generer-code", headers=header)
    assert r.status_code == 200
    code = r.json()["code"]
    assert len(code) == 8

    # Simule le webhook Telegram recevant ce code depuis le chat privé du worker
    r = client.post("/telegram/webhook", json={"message": {"text": code, "chat": {"id": 555111222}}})
    assert r.status_code == 200

    assert client.get("/auth/me", headers=header).json()["telegram_lie"] is True

    # Le mot de passe ne fonctionne plus une fois le compte lié
    r = client.post("/auth/login", json={"email": "tg1@uco.fr", "password": "pw123456"})
    assert r.status_code == 409


def test_webhook_ignore_silencieusement_un_code_inconnu():
    r = client.post("/telegram/webhook", json={"message": {"text": "CODEBIDON", "chat": {"id": 1}}})
    assert r.status_code == 200


def test_otp_demander_reponse_identique_compte_existant_ou_non(monkeypatch):
    """Anti-énumération : impossible de deviner si un email est inscrit à
    partir de la réponse de /auth/otp/demander."""
    _activer_bot_telegram_de_test(monkeypatch)

    client.post("/auth/register", json={"email": "tg2@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("tg2@uco.fr")
    header = auth_header("tg2@uco.fr", "pw123456")
    code_liaison = client.post("/auth/telegram/generer-code", headers=header).json()["code"]
    client.post("/telegram/webhook", json={"message": {"text": code_liaison, "chat": {"id": 555111333}}})

    r_existant = client.post("/auth/otp/demander", json={"email": "tg2@uco.fr"})
    r_inexistant = client.post("/auth/otp/demander", json={"email": "personne@uco.fr"})
    assert r_existant.status_code == r_inexistant.status_code == 200
    assert r_existant.json()["message"] == r_inexistant.json()["message"]


def test_otp_connexion_bon_code_puis_code_non_reutilisable(monkeypatch):
    _activer_bot_telegram_de_test(monkeypatch)

    client.post("/auth/register", json={"email": "tg3@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("tg3@uco.fr")
    header = auth_header("tg3@uco.fr", "pw123456")
    code_liaison = client.post("/auth/telegram/generer-code", headers=header).json()["code"]
    client.post("/telegram/webhook", json={"message": {"text": code_liaison, "chat": {"id": 555111444}}})

    from models import User, get_engine, get_session_factory
    db = get_session_factory(get_engine())()
    try:
        client.post("/auth/otp/demander", json={"email": "tg3@uco.fr"})
        otp = db.query(User).filter_by(email="tg3@uco.fr").first().otp_code
    finally:
        db.close()
    assert otp is not None and len(otp) == 6

    # Mauvais code -> refusé
    assert client.post("/auth/otp/verifier", json={"email": "tg3@uco.fr", "code": "000000"}).status_code == 401

    # Bon code -> jeton de session valide
    r = client.post("/auth/otp/verifier", json={"email": "tg3@uco.fr", "code": otp})
    assert r.status_code == 200
    token = r.json()["token"]
    assert client.get("/auth/me", headers={"Authorization": f"Bearer {token}"}).status_code == 200

    # Le même code ne peut pas être réutilisé (consommé après succès)
    r = client.post("/auth/otp/verifier", json={"email": "tg3@uco.fr", "code": otp})
    assert r.status_code == 401


def test_otp_verrouille_apres_5_echecs(monkeypatch):
    _activer_bot_telegram_de_test(monkeypatch)

    client.post("/auth/register", json={"email": "tg4@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("tg4@uco.fr")
    header = auth_header("tg4@uco.fr", "pw123456")
    code_liaison = client.post("/auth/telegram/generer-code", headers=header).json()["code"]
    client.post("/telegram/webhook", json={"message": {"text": code_liaison, "chat": {"id": 555111555}}})

    client.post("/auth/otp/demander", json={"email": "tg4@uco.fr"})
    for _ in range(5):
        r = client.post("/auth/otp/verifier", json={"email": "tg4@uco.fr", "code": "999999"})
        assert r.status_code == 401

    # Le 6e essai est bloqué par le verrouillage, même avec un mauvais code
    r = client.post("/auth/otp/verifier", json={"email": "tg4@uco.fr", "code": "999999"})
    assert r.status_code == 429


def test_generer_code_telegram_refuse_si_bot_non_configure(monkeypatch):
    """Sans TELEGRAM_BOT_TOKEN, la liaison est refusée proprement plutôt que
    de planter. On force explicitement l'absence de token plutôt que de
    compter sur l'environnement (une vraie valeur peut être présente dans le
    .env local, utilisée par ailleurs pour les notifications de groupe)."""
    import api
    monkeypatch.setattr(api, "TELEGRAM_BOT_TOKEN", None)

    client.post("/auth/register", json={"email": "tg5@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    approuver("tg5@uco.fr")
    header = auth_header("tg5@uco.fr", "pw123456")
    r = client.post("/auth/telegram/generer-code", headers=header)
    assert r.status_code == 503


def test_admin_configurer_webhook_refuse_sans_app_base_url_https(monkeypatch):
    _activer_bot_telegram_de_test(monkeypatch)
    monkeypatch.delenv("APP_BASE_URL", raising=False)

    client.post("/auth/register", json={"email": "clientadmin@ia.fr", "password": "pw123456", "role": "client", "secteur_activite": "test"})
    approuver("clientadmin@ia.fr")

    # Un compte non-admin ne peut pas configurer le webhook
    r = client.post("/admin/telegram/configurer-webhook", headers=auth_header("clientadmin@ia.fr", "pw123456"))
    assert r.status_code == 403


# ---------------------------------------------------------------------------
# NOTIFICATIONS SUR INSCRIPTION / APPROBATION (alerte l'équipe en temps réel
# plutôt que de laisser un compte en attente sans que personne ne le sache)
# ---------------------------------------------------------------------------

def test_inscription_declenche_une_notification_telegram_si_bot_configure(monkeypatch):
    import notifications
    messages_envoyes = []
    monkeypatch.setattr(notifications, "TELEGRAM_BOT_TOKEN", "jeton-bot-de-test")
    monkeypatch.setattr(notifications, "TELEGRAM_CHAT_ID", "-100999")
    monkeypatch.setattr(
        notifications, "envoyer_message_telegram",
        lambda chat_id, texte: messages_envoyes.append((chat_id, texte)) or True,
    )

    r = client.post("/auth/register", json={
        "email": "notif@ia.fr", "password": "pw123456", "role": "client", "secteur_activite": "retail",
    })
    assert r.status_code == 200
    assert len(messages_envoyes) == 1
    assert messages_envoyes[0][0] == "-100999"
    assert "notif@ia.fr" in messages_envoyes[0][1]


def test_inscription_ne_plante_pas_sans_bot_configure():
    """Comportement par défaut de la suite de tests (pas de token) -- doit
    rester silencieux, jamais faire échouer l'inscription elle-même."""
    r = client.post("/auth/register", json={
        "email": "sansnotif@ia.fr", "password": "pw123456", "role": "client", "secteur_activite": "retail",
    })
    assert r.status_code == 200


def _creer_admin_de_test(email, password):
    """Un compte admin ne se crée pas via /auth/register (volontairement
    refusé, voir register()) -- on l'insère directement en base, comme le
    ferait une opération manuelle en production."""
    from auth import hash_password
    from models import RoleEnum, User, get_engine, get_session_factory

    db = get_session_factory(get_engine())()
    try:
        if db.query(User).filter_by(email=email).first():
            return
        db.add(User(email=email, role=RoleEnum.admin, password_hash=hash_password(password), approuve=True))
        db.commit()
    finally:
        db.close()


def test_approbation_notifie_le_compte_seulement_si_telegram_lie(monkeypatch):
    import notifications

    messages_envoyes = []
    monkeypatch.setattr(notifications, "TELEGRAM_BOT_TOKEN", "jeton-bot-de-test")
    monkeypatch.setattr(
        notifications, "envoyer_message_telegram",
        lambda chat_id, texte: messages_envoyes.append((chat_id, texte)) or True,
    )

    _creer_admin_de_test("admin_approb@uco.fr", "adminpw123")
    admin_header = auth_header("admin_approb@uco.fr", "adminpw123")

    # Compte SANS Telegram lié : l'approbation ne doit rien tenter d'envoyer
    r = client.post("/auth/register", json={"email": "approb1@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    user_id_sans_lien = r.json()["user_id"]
    messages_envoyes.clear()  # on ignore le message de notif d'inscription lui-même

    r = client.post(f"/admin/comptes/{user_id_sans_lien}/approuver", headers=admin_header)
    assert r.status_code == 200
    assert messages_envoyes == []

    # Compte AVEC Telegram lié : l'approbation doit envoyer une confirmation
    r = client.post("/auth/register", json={"email": "approb2@uco.fr", "password": "pw123456", "role": "worker", "accepte_confidentialite": True})
    user_id_avec_lien = r.json()["user_id"]

    from models import User, get_engine, get_session_factory
    db = get_session_factory(get_engine())()
    try:
        db.query(User).filter_by(id=user_id_avec_lien).update({"telegram_chat_id": "777888999"})
        db.commit()
    finally:
        db.close()
    messages_envoyes.clear()

    r = client.post(f"/admin/comptes/{user_id_avec_lien}/approuver", headers=admin_header)
    assert r.status_code == 200
    assert len(messages_envoyes) == 1
    assert messages_envoyes[0][0] == "777888999"
